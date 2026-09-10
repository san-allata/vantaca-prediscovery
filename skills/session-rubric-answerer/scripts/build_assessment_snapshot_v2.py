#!/usr/bin/env python3
"""
build_assessment_snapshot_v2.py  —  TEST COPY (v2)

Writes the Assessment sheet xlsx from a list of fully merged,
answered, and classified rubric rows.

Input (input.json):
  {
    "rows": [
      {
        "uid":               "Q0001",
        "qid":               "1.1.1",
        "rubric_row":        5,
        "domain":            "Platform Foundation",
        "cap_num":           1.1,
        "capability":        "General System Access",
        "dimension":         "Multi-system navigation",
        "priority":          "P0",
        "discovery_question": "How many...",
        "townsq_capability": "After config — ...",
        "branch_answer":     "Branch uses 4 systems...",
        "classification":    "AC",
        "fb_priority":       "",
        "assessor_notes":    "No digital work order tracking today."
      },
      ...
    ]
  }

Output:
  - Assessment_v2.xlsx written to working directory
  - stdout JSON: { "output_filename": "Assessment_v2.xlsx",
                   "rows_written": N,
                   "sheet_name": "Assessment",
                   "_source_version": "v2" }

Fast-fail: exits immediately with rows_written=0 if rows is
missing, null, or empty. No hanging.
"""

import json
import os
import sys
from datetime import datetime, timezone

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:
    print(json.dumps({"error": "openpyxl not available", "rows_written": 0}))
    sys.exit(1)


# ---------------------------------------------------------------------------
# Colour palette
# ---------------------------------------------------------------------------
COLOR_HEADER      = "1F3864"
COLOR_HEADER_FONT = "FFFFFF"
COLOR_PC          = "FFF2CC"
COLOR_AC          = "D9EAD3"
COLOR_FB          = "FCE5CD"
COLOR_NA          = "F3F3F3"
COLOR_ROW_ALT     = "EAF4FB"

COLUMNS = [
    ("Dom #",              "dom_num",            8),
    ("Domain",             "domain",             22),
    ("Cap #",              "cap_num",            8),
    ("Capability",         "capability",         24),
    ("Dimension",          "dimension",          24),
    ("Priority",           "priority",           10),
    ("Discovery Question", "discovery_question", 50),
    ("Branch Answer",      "branch_answer",      55),
    ("TownSq Capability",  "townsq_capability",  50),
    ("Classification",     "classification",     16),
    ("FB Priority",        "fb_priority",        12),
    ("Assessor Notes",     "assessor_notes",     45),
    ("QID",                "qid",                12),
    ("UID",                "uid",                12),
]


def class_fill(classification: str):
    c = str(classification).strip().upper()
    mapping = {"PC": COLOR_PC, "AC": COLOR_AC, "FB": COLOR_FB, "NA": COLOR_NA}
    return PatternFill("solid", fgColor=mapping[c]) if c in mapping else None


def thin_border():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)


def build_worksheet(ws, rows: list):
    header_font  = Font(bold=True, color=COLOR_HEADER_FONT, size=10)
    header_fill  = PatternFill("solid", fgColor=COLOR_HEADER)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx, (label, _, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font      = header_font
        cell.fill      = header_fill
        cell.alignment = header_align
        cell.border    = thin_border()
        ws.column_dimensions[
            openpyxl.utils.get_column_letter(col_idx)
        ].width = width

    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    data_align  = Alignment(vertical="top", wrap_text=True)
    alt_fill    = PatternFill("solid", fgColor=COLOR_ROW_ALT)
    normal_fill = PatternFill("solid", fgColor="FFFFFF")

    for row_idx, row in enumerate(rows, start=2):
        cl   = str(row.get("classification", "")).strip().upper()
        fill = class_fill(cl) or (alt_fill if row_idx % 2 == 0 else normal_fill)

        for col_idx, (_, field_key, _) in enumerate(COLUMNS, start=1):
            val = row.get(field_key, "") or ""
            cell = ws.cell(row=row_idx, column=col_idx, value=str(val))
            cell.alignment = data_align
            cell.fill      = fill
            cell.border    = thin_border()

        ws.row_dimensions[row_idx].height = 60


def main():
    # --- Load input.json ---
    input_path = os.path.join(os.getcwd(), "input.json")

    # Fast-fail: no input file
    if not os.path.exists(input_path):
        print(json.dumps({
            "output_filename": None,
            "rows_written":    0,
            "sheet_name":      "Assessment",
            "_source_version": "v2",
            "warning":         "input.json not found — nothing to write",
        }))
        return

    with open(input_path, "r", encoding="utf-8") as f:
        try:
            payload = json.load(f)
        except json.JSONDecodeError as e:
            print(json.dumps({
                "output_filename": None,
                "rows_written":    0,
                "_source_version": "v2",
                "error":           f"input.json parse error: {e}",
            }))
            sys.exit(1)

    rows = payload.get("rows") or []

    # Fast-fail: empty or missing rows
    if not rows:
        print(json.dumps({
            "output_filename": None,
            "rows_written":    0,
            "sheet_name":      "Assessment",
            "_source_version": "v2",
            "warning":         "rows array is empty or missing — nothing to write",
        }))
        return

    # Derive dom_num from cap_num if not present
    for r in rows:
        if not r.get("dom_num"):
            try:
                r["dom_num"] = int(float(str(r.get("cap_num", 0)).split(".")[0]))
            except (ValueError, TypeError):
                r["dom_num"] = ""

    # --- Build workbook ---
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Assessment"

    build_worksheet(ws, rows)

    meta_ws = wb.create_sheet("_meta")
    meta_ws["A1"] = "Generated by"
    meta_ws["B1"] = "build_assessment_snapshot_v2.py"
    meta_ws["A2"] = "Timestamp (UTC)"
    meta_ws["B2"] = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    meta_ws["A3"] = "Rows written"
    meta_ws["B3"] = len(rows)
    meta_ws["A4"] = "Source version"
    meta_ws["B4"] = "v2 — JSON rubric index"

    output_filename = "Assessment_v2.xlsx"
    wb.save(output_filename)

    print(json.dumps({
        "output_filename": output_filename,
        "rows_written":    len(rows),
        "sheet_name":      "Assessment",
        "_source_version": "v2",
    }))


if __name__ == "__main__":
    main()
