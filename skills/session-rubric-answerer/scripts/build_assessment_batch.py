import base64
import io
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Load input
# ---------------------------------------------------------------------------
with open("input.json") as f:
    data = json.load(f)

rows = data.get("rows", [])
branch = str(data.get("branch", "Branch")).strip()
session = str(data.get("session", "0")).strip()
run = str(data.get("run", "1")).strip()

if not rows:
    print("ERROR: no rows in input", file=sys.stderr)
    sys.exit(1)

# ---------------------------------------------------------------------------
# Style helpers
# ---------------------------------------------------------------------------
def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def font(bold=True, color="000000", size=10):
    return Font(bold=bold, color=color, size=size)

def wrap_top():
    return Alignment(wrap_text=True, vertical="top")

def wrap_center():
    return Alignment(wrap_text=True, vertical="center", horizontal="center")

HEADER_FILL = fill("1F4E79")
HEADER_FONT = font(bold=True, color="FFFFFF", size=11)

# Classification fills - HITL rows get amber fill; cell value stays blank
CLASS_FILLS = {
    "AC": fill("70AD47"),
    "FB": fill("FF0000"),
    "PC": fill("4472C4"),
    "NA": fill("FFD966"),
    "": fill("FFC000"),  # amber for blank/HITL rows
}
CLASS_FONTS = {
    "AC": font(bold=True, color="FFFFFF"),
    "FB": font(bold=True, color="FFFFFF"),
    "PC": font(bold=True, color="FFFFFF"),
    "NA": font(bold=True, color="000000"),
    "": font(bold=True, color="000000"),
}

# Proximity fills
PROX_FILLS = {
    "Exact match":     fill("70AD47"),
    "High (~75%)": fill("A9D18E"),
    "Moderate (~50%)": fill("FFD966"),
    "Low (~25%)": fill("F4B942"),
    "No match":        fill("FF0000"),
    "": fill("FFC000"),
}
PROX_FONTS = {
    "Exact match":     font(bold=True, color="FFFFFF"),
    "High (~75%)": font(bold=True, color="000000"),
    "Moderate (~50%)": font(bold=True, color="000000"),
    "Low (~25%)": font(bold=True, color="000000"),
    "No match":        font(bold=True, color="FFFFFF"),
    "": font(bold=True, color="000000"),
}

# Column widths for 15 columns
COL_WIDTHS = [10, 10, 8, 14, 18, 8, 22, 28, 9, 45, 55, 45, 14, 60, 18]

HEADERS = [
    "UID", "QID", "Session", "Planned Question ID",
    "Domain", "Cap #", "Capability", "Dimension", "Priority",
    "Discovery Question", "Branch Answer", "TownSq Capability",
    "Classification", "Assessor Notes", "Proximity"
]

# ---------------------------------------------------------------------------
# Build workbook
# ---------------------------------------------------------------------------
wb = Workbook()
ws = wb.active
ws.title = "Assessment"

# Header row
for col_idx, header in enumerate(HEADERS, start=1):
    cell = ws.cell(row=1, column=col_idx, value=header)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = wrap_center()

# Column widths
for col_idx, width in enumerate(COL_WIDTHS, start=1):
    ws.column_dimensions[get_column_letter(col_idx)].width = width

ws.row_dimensions[1].height = 30
ws.freeze_panes = "A2"

# ---------------------------------------------------------------------------
# Data rows
# ---------------------------------------------------------------------------
for row_idx, row in enumerate(rows, start=2):

    # --- Resolve classification ---
    raw_class = str(row.get("classification") or "").strip().upper()

    # Strip "HITL" - it is never a valid cell value
    if raw_class == "HITL":
        raw_class = ""

    # Force blank for any value outside the allowed set
    if raw_class not in ("PC", "AC", "FB", "NA", ""):
        raw_class = ""

    # A row is HITL if hitl flag is set, or classification is blank
    is_hitl = bool(row.get("hitl", False)) or raw_class == ""

    # If HITL, classification cell must be blank
    classification = "" if is_hitl else raw_class

    # --- Proximity ---
    proximity = str(row.get("proximity") or "").strip()
    if is_hitl:
        proximity = ""

    # --- Write all 15 columns ---
    values = [
        row.get("uid", ""),
        row.get("qid", ""),
        row.get("session", ""),
        row.get("planned_question_id", ""),
        row.get("domain", ""),
        row.get("cap_num", ""),
        row.get("capability", ""),
        row.get("dimension", ""),
        row.get("priority", ""),
        row.get("discovery_question", ""),
        row.get("branch_answer", ""),
        row.get("townsq_capability", ""),
        classification,  # col 13 - blank for HITL, code for classified
        row.get("assessor_notes", ""),
        proximity,       # col 15
    ]

    for col_idx, value in enumerate(values, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.alignment = wrap_top()

    # --- Style col 13: Classification ---
    class_cell = ws.cell(row=row_idx, column=13)
    class_key = "" if is_hitl else classification
    class_cell.fill = CLASS_FILLS.get(class_key, CLASS_FILLS[""])
    class_cell.font = CLASS_FONTS.get(class_key, CLASS_FONTS[""])

    # --- Style col 15: Proximity ---
    prox_cell = ws.cell(row=row_idx, column=15)
    prox_key = proximity if proximity in PROX_FILLS else ""
    prox_cell.fill = PROX_FILLS[prox_key]
    prox_cell.font = PROX_FONTS[prox_key]

    ws.row_dimensions[row_idx].height = 80

# ---------------------------------------------------------------------------
# Always output as base64 to stdout
# ---------------------------------------------------------------------------
buf = io.BytesIO()
wb.save(buf)
buf.seek(0)
print(base64.b64encode(buf.read()).decode("utf-8"))
