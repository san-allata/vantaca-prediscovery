"""xlsx_builder.py — Assessment xlsx builder for Session Rubric Answerer.

Usage (via execute_code):
  - Read this file via read_skill_resource("references/xlsx_builder.py")
  - Embed the returned content verbatim as the execute_code `code` parameter
  - Pass row data via the execute_code `input` parameter (written as input.json)
  - The script reads input.json, builds the xlsx, writes it to disk
  - execute_code captures it in outputFiles — use that blobKey for render_content

Input JSON schema:
  {
    "branch": "CAMS",
    "session": 6,
    "run": 1,
    "rows": [
      {
        "uid": "Q0285",
        "qid": "6.8.1",
        "session": 6,
        "planned_question_id": "S6-Q12",
        "domain": "Financial Operations",
        "cap_num": "6.8",
        "capability": "Supply Charges",
        "dimension": "Ancillary billing visibility",
        "priority": "P1",
        "discovery_question": "How do you view...",
        "branch_answer": "We maintain...",
        "townsq_capability": "Ancillary charges...",
        "classification": "AC",
        "assessor_notes": "BRANCH: ...",
        "proximity": "High (~75%)",
        "hitl": false
      }
    ]
  }

Output: Assessment_{branch}_Session{N}_Run{M}.xlsx written to current directory.
"""

import json
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ── Colour palette ────────────────────────────────────────────────────────────
COL_HEADER = "1F4E79"   # dark blue header
COL_PC     = "FFF2CC"   # yellow  — Process Change
COL_AC     = "E2EFDA"   # green   — Adoption/Config
COL_FB     = "FCE4D6"   # orange  — Feature Backlog
COL_NA     = "EDEDED"   # grey    — Not Applicable
COL_HITL   = "FFEB9C"   # amber   — HITL / blank classification
COL_WHITE  = "FFFFFF"

FILL = {
    "PC": PatternFill("solid", fgColor=COL_PC),
    "AC": PatternFill("solid", fgColor=COL_AC),
    "FB": PatternFill("solid", fgColor=COL_FB),
    "NA": PatternFill("solid", fgColor=COL_NA),
    "":   PatternFill("solid", fgColor=COL_HITL),
}

HEADER_FILL  = PatternFill("solid", fgColor=COL_HEADER)
HEADER_FONT  = Font(bold=True, color=COL_WHITE, size=10, name="Calibri")
DATA_FONT    = Font(size=9, name="Calibri")
WRAP         = Alignment(wrap_text=True, vertical="top")
HEADER_ALIGN = Alignment(wrap_text=True, vertical="center", horizontal="center")


# ── Column definitions ────────────────────────────────────────────────────────
# (header, field_key, width)
COLUMNS = [
    ("UID",                   "uid",                  10),
    ("QID",                   "qid",                   8),
    ("Session",               "session",               8),
    ("Planned Question ID",   "planned_question_id",  12),
    ("Domain",                "domain",               18),
    ("Cap #",                 "cap_num",               7),
    ("Capability",            "capability",           22),
    ("Dimension",             "dimension",            24),
    ("Priority",              "priority",              8),
    ("Discovery Question",    "discovery_question",   45),
    ("Branch Answer",         "branch_answer",        55),
    ("TownSq Capability",     "townsq_capability",    55),
    ("Classification",        "classification",       14),
    ("Assessor Notes",        "assessor_notes",       55),
    ("Proximity",             "proximity",            16),
]


def _thin_border():
    side = Side(style="thin", color="BFBFBF")
    return Border(left=side, right=side, top=side, bottom=side)


def build(data: dict) -> str:
    """Build the Assessment xlsx from a data dict. Returns the output filename."""
    branch  = data.get("branch", "BRANCH")
    session = data.get("session", 0)
    run     = data.get("run", 1)
    rows    = data.get("rows", [])

    filename = f"Assessment_{branch}_Session{session}_Run{run}.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Assessment"

    # ── Header row ────────────────────────────────────────────────────────────
    for col_idx, (header, _, _width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill      = HEADER_FILL
        cell.font      = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        cell.border    = _thin_border()

    ws.row_dimensions[1].height = 30

    # ── Data rows ─────────────────────────────────────────────────────────────
    for row_idx, row in enumerate(rows, 2):
        classification = str(row.get("classification", "")).strip().upper()
        # Only valid codes pass through; everything else → HITL amber
        if classification not in ("PC", "AC", "FB", "NA"):
            classification = ""

        row_fill = FILL.get(classification, FILL[""])

        for col_idx, (_, field_key, _) in enumerate(COLUMNS, 1):
            raw_value = row.get(field_key, "")

            # Classification column: write normalised code (blank for HITL)
            if field_key == "classification":
                value = classification
            else:
                value = raw_value if raw_value is not None else ""

            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill      = row_fill
            cell.font      = DATA_FONT
            cell.alignment = WRAP
            cell.border    = _thin_border()

        ws.row_dimensions[row_idx].height = 60

    # ── Column widths ─────────────────────────────────────────────────────────
    for col_idx, (_, _, width) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # ── Freeze header + auto-filter ───────────────────────────────────────────
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"

    wb.save(filename)
    print(f"OK: {filename} — {len(rows)} rows written.")
    return filename


# ── Entry point (called when run as __main__ inside execute_code) ─────────────
if __name__ == "__main__":
    with open("input.json", "r", encoding="utf-8") as f:
        data = json.load(f)
    build(data)
